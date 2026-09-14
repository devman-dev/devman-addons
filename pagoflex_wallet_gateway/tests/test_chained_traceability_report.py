import hashlib
from io import BytesIO

from openpyxl import load_workbook

from odoo.exceptions import AccessError, ValidationError
from odoo.tests import TransactionCase, tagged


@tagged("post_install", "-at_install")
class TestPfGatewayChainedTraceabilityReport(TransactionCase):
    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.User = cls.env["pf.gateway.user"]
        cls.Account = cls.env["pf.gateway.bank.account"]
        cls.Transfer = cls.env["pf.gateway.transfer"]
        cls.Wizard = cls.env["pf.gateway.chained.traceability.report.wizard"]
        cls.Audit = cls.env["pf.gateway.traceability.report.audit"]

        cls.user_a = cls.User.create(
            {"external_id": "trace-user-a", "full_name": "A", "cuit_cuil": "20-12345678-6"}
        )
        cls.user_b = cls.User.create(
            {"external_id": "trace-user-b", "full_name": "B", "cuit_cuil": "20-23456789-3"}
        )
        cls.user_c = cls.User.create(
            {"external_id": "trace-user-c", "full_name": "C", "cuit_cuil": "20-00000000-1"}
        )
        cls.account_a = cls.Account.create(
            {
                "external_id": "trace-account-a", "gateway_user_id": cls.user_a.id,
                "cvu_cbu": "0000003100000000000001", "status": "active", "app": "pagoflex",
            }
        )
        cls.account_b = cls.Account.create(
            {
                "external_id": "trace-account-b", "gateway_user_id": cls.user_b.id,
                "cvu_cbu": "0000003100000000000002", "status": "blocked", "app": "pagoflex",
            }
        )
        cls.Transfer.create(
            {
                "external_id": "trace-transfer-internal", "movement_nature": "TRANSFER",
                "status": "COMPLETED", "amount": 100, "transaction_at": "2026-01-10 12:00:00",
                "source_user_id": cls.user_a.id, "destination_user_id": cls.user_b.id,
                "source_bank_account_id": cls.account_a.id,
                "destination_bank_account_id": cls.account_b.id,
                "source_address": cls.account_a.cvu_cbu, "destination_address": cls.account_b.cvu_cbu,
            }
        )
        cls.Transfer.create(
            {
                "external_id": "trace-transfer-to-cuit-only", "movement_nature": "TRANSFER",
                "status": "COMPLETED", "amount": 25, "transaction_at": "2026-01-12 12:00:00",
                "source_user_id": cls.user_a.id, "source_bank_account_id": cls.account_a.id,
                "source_address": cls.account_a.cvu_cbu,
                "destination_address": "2850590940090418135299",
                "destination_owner_id_type": "CUIT", "destination_owner_id": "20-00000000-1",
                "destination_owner_name": "Usuario C",
            }
        )
        cls.Transfer.create(
            {
                "external_id": "trace-transfer-from-cuit-only", "movement_nature": "TRANSFER",
                "status": "COMPLETED", "amount": 10, "transaction_at": "2026-01-13 12:00:00",
                "source_user_id": cls.user_c.id, "source_address": "2850590940090418135299",
                "destination_address": "2850590940090418135300", "destination_owner_name": "Externo 2",
            }
        )
        cls.Transfer.create(
            {
                "external_id": "trace-transfer-external", "movement_nature": "TRANSFER",
                "status": "COMPLETED", "amount": 75, "transaction_at": "2026-01-11 12:00:00",
                "source_user_id": cls.user_b.id, "source_bank_account_id": cls.account_b.id,
                "source_address": cls.account_b.cvu_cbu,
                "destination_address": "2850590940090418135201", "destination_owner_name": "Externo",
            }
        )

    def _audit(self, scope="both"):
        return self.Audit.sudo().create_request(
            {
                "cuits": ["20123456786"], "all_cuits": False,
                "date_from": "2026-01-01", "date_to": "2026-01-31",
                "timezone": "America/Argentina/Buenos_Aires", "operation_scope": scope, "max_depth": 5,
            }
        )[0]

    def test_cuit_validation(self):
        wizard = self.Wizard.new({"cuit_text": "20-12345678-6"})
        self.assertEqual(wizard._normalized_cuits(), ["20123456786"])
        wizard.cuit_text = "20-12345678-9"
        with self.assertRaises(ValidationError):
            wizard._normalized_cuits()

    def test_bank_account_normalizes_blocked_status(self):
        self.assertEqual(self.account_b._normalize_status("BLOCKED"), "blocked")

    def test_external_filter_traverses_internal_edge(self):
        content, count = self.env["pf.gateway.chained.traceability.report.service"].generate_xlsx(
            self._audit("external")
        )
        self.assertEqual(count, 3)
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
        rows = list(workbook["Trazabilidad"].iter_rows(values_only=True))
        headers = list(rows[0])
        payloads = {row[4]: dict(zip(headers, row)) for row in rows[1:]}
        payload = payloads["trace-transfer-external"]
        self.assertEqual(payload["nivel"], "2")
        self.assertEqual(payload["estado_cvu_origen"], "BLOQUEADO")
        self.assertEqual(payload["estado_cvu_destino"], "NO_APLICA_EXTERNO")
        self.assertEqual(payloads["trace-transfer-to-cuit-only"]["nivel"], "1")
        self.assertEqual(payloads["trace-transfer-from-cuit-only"]["nivel"], "2")

    def test_internal_filter_exports_only_internal_edge(self):
        content, count = self.env["pf.gateway.chained.traceability.report.service"].generate_xlsx(
            self._audit("internal")
        )
        self.assertEqual(count, 1)
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
        rows = list(workbook["Trazabilidad"].iter_rows(values_only=True))
        payload = dict(zip(rows[0], rows[1]))
        self.assertEqual(payload["external_id"], "trace-transfer-internal")
        self.assertEqual(payload["estado_cvu_destino"], "BLOQUEADO")

    def test_chain_continues_from_exact_destination_cvu(self):
        account_b2 = self.Account.create(
            {
                "external_id": "trace-account-b2", "gateway_user_id": self.user_b.id,
                "cvu_cbu": "0000003100000000000003", "status": "active", "app": "pagoflex",
            }
        )
        movement_to_b2 = self.Transfer.create(
            {
                "external_id": "trace-transfer-to-b2", "movement_nature": "TRANSFER",
                "status": "COMPLETED", "amount": 20, "transaction_at": "2026-01-14 12:00:00",
                "source_user_id": self.user_b.id, "destination_user_id": self.user_b.id,
                "source_bank_account_id": self.account_b.id, "destination_bank_account_id": account_b2.id,
                "source_address": self.account_b.cvu_cbu, "destination_address": account_b2.cvu_cbu,
            }
        )
        movement_from_b2 = self.Transfer.create(
            {
                "external_id": "trace-transfer-from-b2", "movement_nature": "TRANSFER",
                "status": "COMPLETED", "amount": 15, "transaction_at": "2026-01-15 12:00:00",
                "source_user_id": self.user_b.id, "source_bank_account_id": account_b2.id,
                "source_address": account_b2.cvu_cbu, "destination_address": "2850590940090418135310",
            }
        )
        content, unused_count = self.env["pf.gateway.chained.traceability.report.service"].generate_xlsx(
            self._audit("both")
        )
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
        rows = list(workbook["Trazabilidad"].iter_rows(values_only=True))
        headers = list(rows[0])
        payloads = {row[4]: dict(zip(headers, row)) for row in rows[1:]}
        self.assertEqual(payloads["trace-transfer-to-b2"]["nivel"], "2")
        self.assertEqual(payloads["trace-transfer-from-b2"]["nivel"], "3")
        self.assertEqual(
            payloads["trace-transfer-from-b2"]["ruta_transferencias"],
            "%s>%s>%s" % (
                self.Transfer.search([("external_id", "=", "trace-transfer-internal")]).id,
                movement_to_b2.id,
                movement_from_b2.id,
            ),
        )

    def test_all_cuits_exports_each_transfer_once_without_recursive_duplicates(self):
        audit = self.Audit.sudo().create_request(
            {
                "cuits": [], "all_cuits": True,
                "date_from": "2026-01-01", "date_to": "2026-01-31",
                "timezone": "America/Argentina/Buenos_Aires", "operation_scope": "both", "max_depth": 10,
            }
        )[0]
        content, count = self.env["pf.gateway.chained.traceability.report.service"].generate_xlsx(audit)
        self.assertEqual(count, 4)
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
        rows = list(workbook["Trazabilidad"].iter_rows(values_only=True))
        self.assertEqual(
            {row[4] for row in rows[1:]},
            {
                "trace-transfer-internal", "trace-transfer-external",
                "trace-transfer-to-cuit-only", "trace-transfer-from-cuit-only",
            },
        )
        self.assertEqual({row[1] for row in rows[1:]}, {"1"})

    def test_audit_is_immutable_and_parameter_hash_is_present(self):
        audit = self._audit()
        self.assertEqual(len(audit.parameter_hash), 64)
        self.assertEqual(audit.parameter_hash, hashlib.sha256(
            json_canonical(audit).encode()
        ).hexdigest())
        with self.assertRaises(AccessError):
            audit.write({"max_depth": 2})
        with self.assertRaises(AccessError):
            audit.unlink()


def json_canonical(audit):
    import json
    return json.dumps(
        {
            "all_cuits": False, "cuits": ["20123456786"], "date_from": "2026-01-01",
            "date_to": "2026-01-31", "timezone": "America/Argentina/Buenos_Aires",
            "operation_scope": audit.operation_scope, "max_depth": 5,
            "export_format": "xlsx", "schema_version": "1.0",
        },
        ensure_ascii=True, sort_keys=True, separators=(",", ":"),
    )
