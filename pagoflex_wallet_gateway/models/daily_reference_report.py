import base64
import json
import re
import unicodedata
from datetime import date, datetime
from io import BytesIO

from openpyxl import load_workbook

from odoo import _, api, fields, models
from odoo.exceptions import UserError


def _normalize_header(value):
    text = (value or "").strip()
    # quitar BOM y caracteres de control invisibles que pueden estar en la primera celda
    text = text.lstrip("\ufeff\u200b\u200c\u200d\xa0")
    text = text.strip().lower()
    if not text:
        return ""
    text = unicodedata.normalize("NFKD", text)
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    text = re.sub(r"[^a-z0-9]+", " ", text).strip()
    return text


def _to_text(value):
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(value, date):
        return value.strftime("%Y-%m-%d")
    return str(value).strip()


def _to_date(value):
    if not value:
        return False
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = _to_text(value)
    if not text:
        return False
    # openpyxl devuelve datetime → _to_text produce "YYYY-MM-DD HH:MM:SS"; quitar la parte horaria
    if " " in text:
        text = text.split(" ")[0]
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%m/%d/%Y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return False


def _to_float(value):
    if value in (None, ""):
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace(" ", "")
    if not text:
        return 0.0
    if "," in text and "." in text:
        text = text.replace(".", "").replace(",", ".")
    else:
        text = text.replace(",", ".")
    try:
        return float(text)
    except ValueError:
        return 0.0


class PfGatewayDailyReferenceImport(models.Model):
    _name = "pf.gateway.daily.reference.import"
    _description = "Importacion de Listado de Referencia Diario"
    _order = "id desc"

    name = fields.Char(string="Nombre", required=True)
    company_id = fields.Many2one(
        "res.company",
        string="Compania",
        default=lambda self: self.env.company,
        required=True,
        readonly=True,
    )
    import_date = fields.Date(string="Fecha de reporte")
    imported_at = fields.Datetime(string="Importado el", default=fields.Datetime.now, readonly=True)
    imported_by_id = fields.Many2one(
        "res.users",
        string="Importado por",
        default=lambda self: self.env.user,
        readonly=True,
    )
    source_file_name = fields.Char(string="Archivo")
    source_headers_json = fields.Text(string="Cabeceras origen")
    line_ids = fields.One2many(
        "pf.gateway.daily.reference.import.line",
        "import_id",
        string="Lineas",
        readonly=True,
    )
    line_count = fields.Integer(string="Cantidad de lineas", compute="_compute_line_count")

    @api.depends("line_ids")
    def _compute_line_count(self):
        for rec in self:
            rec.line_count = len(rec.line_ids)

    def action_open_import_wizard(self):
        self.ensure_one()
        action = self.env["ir.actions.actions"]._for_xml_id(
            "pagoflex_wallet_gateway.action_pf_gateway_daily_reference_import_wizard"
        )
        action["context"] = dict(self.env.context, default_import_date=self.import_date)
        return action


class PfGatewayDailyReferenceImportLine(models.Model):
    _name = "pf.gateway.daily.reference.import.line"
    _description = "Linea importada de Listado de Referencia Diario"
    _order = "import_id desc, sequence asc, id asc"

    import_id = fields.Many2one(
        "pf.gateway.daily.reference.import",
        string="Importacion",
        required=True,
        ondelete="cascade",
    )
    sequence = fields.Integer(string="Fila")
    # --- columnas principales del reporte ---
    tipo_archivo = fields.Char(string="Tipo archivo")
    operation_date = fields.Date(string="Fecha negocio")
    nro_registro = fields.Char(string="Nro registro")
    tipo_obj = fields.Char(string="Tipo obj")
    id_debin = fields.Char(string="ID DEBIN")
    tipo_mov = fields.Char(string="Tipo mov")
    fecha_hora_coelsa = fields.Char(string="Fecha/hora COELSA")
    origen_trx = fields.Char(string="Origen TRX")
    moneda = fields.Char(string="Moneda")
    cbu = fields.Char(string="CBU")
    denominacion = fields.Char(string="Denominacion")
    cuit = fields.Char(string="CUIT")
    amount = fields.Float(string="Importe COELSA", digits=(16, 2))
    estado = fields.Char(string="Estado COELSA")
    concepto = fields.Char(string="Concepto")
    id_movimiento = fields.Char(string="ID Movimiento")
    detalle = fields.Char(string="Detalle")
    importe_2 = fields.Float(string="Importe 2", digits=(16, 2))
    cvu = fields.Char(string="CVU")
    cuit_virtual = fields.Char(string="CUIT virtual")
    cvu_credito = fields.Char(string="CVU credito")
    mismo_tit = fields.Char(string="Mismo titular")
    credito_forzado = fields.Char(string="Credito forzado")
    tipo_reverso = fields.Char(string="Tipo reverso")
    raw_payload = fields.Text(string="Fila origen")


class PfGatewayDailyReferenceImportWizard(models.TransientModel):
    _name = "pf.gateway.daily.reference.import.wizard"
    _description = "Asistente de importacion de Listado de Referencia Diario"

    file_data = fields.Binary(string="Archivo XLSX", required=True)
    file_name = fields.Char(string="Nombre de archivo")
    import_date = fields.Date(string="Fecha de reporte", default=fields.Date.context_today)
    has_header = fields.Boolean(string="El archivo tiene cabecera", default=True)
    sheet_name = fields.Char(string="Hoja", help="Opcional. Si no se indica, se usa la hoja activa.")

    def action_import_file(self):
        self.ensure_one()
        if not self.file_data:
            raise UserError(_("Debe adjuntar un archivo XLSX."))

        try:
            workbook = load_workbook(filename=BytesIO(base64.b64decode(self.file_data)), data_only=True)
        except Exception as exc:
            raise UserError(_("No se pudo leer el archivo XLSX: %s") % exc) from exc

        if self.sheet_name:
            if self.sheet_name not in workbook.sheetnames:
                raise UserError(_("La hoja '%s' no existe en el archivo.") % self.sheet_name)
            sheet = workbook[self.sheet_name]
        else:
            sheet = workbook.active

        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            raise UserError(_("La hoja seleccionada no contiene datos."))

        first_row = list(rows[0] or [])
        if self.has_header:
            headers = [_normalize_header(_to_text(val)) or f"col_{idx + 1}" for idx, val in enumerate(first_row)]
            data_rows = rows[1:]
        else:
            headers = [f"col_{idx + 1}" for idx in range(len(first_row))]
            data_rows = rows

        import_record = self.env["pf.gateway.daily.reference.import"].create(
            {
                "name": _("Listado de Referencia Diario %s") % fields.Datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                "import_date": self.import_date,
                "source_file_name": self.file_name,
                "source_headers_json": json.dumps(headers, ensure_ascii=True),
            }
        )

        def _col(payload, name):
            return payload.get(name, "") or ""

        line_values = []
        for row_idx, row in enumerate(data_rows, start=2 if self.has_header else 1):
            row_list = list(row or [])
            if not any(value not in (None, "") for value in row_list):
                continue

            payload = {}
            for col_idx, raw_value in enumerate(row_list):
                key = headers[col_idx] if col_idx < len(headers) else f"col_{col_idx + 1}"
                payload[key] = _to_text(raw_value)

            line_values.append(
                {
                    "import_id": import_record.id,
                    "sequence": row_idx,
                    "tipo_archivo": _col(payload, "tipo archivo"),
                    "operation_date": _to_date(_col(payload, "fecha negocio")),
                    "nro_registro": _col(payload, "nro registro"),
                    "tipo_obj": _col(payload, "tipo obj"),
                    "id_debin": _col(payload, "id debin"),
                    "tipo_mov": _col(payload, "tipo mov coelsa"),
                    "fecha_hora_coelsa": _col(payload, "fecha hora coelsa"),
                    "origen_trx": _col(payload, "origen trx"),
                    "moneda": _col(payload, "moneda"),
                    "cbu": _col(payload, "cbu"),
                    "denominacion": _col(payload, "denominacion"),
                    "cuit": _col(payload, "cuit"),
                    "amount": _to_float(_col(payload, "importe coelsa")),
                    "estado": _col(payload, "estado coelsa"),
                    "concepto": _col(payload, "concepto"),
                    "id_movimiento": _col(payload, "id movimiento"),
                    "detalle": _col(payload, "detalle"),
                    "importe_2": _to_float(_col(payload, "importe 2")),
                    "cvu": _col(payload, "cvu"),
                    "cuit_virtual": _col(payload, "cuit virtual"),
                    "cvu_credito": _col(payload, "cvu credito"),
                    "mismo_tit": _col(payload, "mismo tit"),
                    "credito_forzado": _col(payload, "credito forzado"),
                    "tipo_reverso": _col(payload, "tipo reverso"),
                    "raw_payload": json.dumps(payload, ensure_ascii=True),
                }
            )

        if not line_values:
            raise UserError(_("No se encontraron filas de datos para importar."))

        self.env["pf.gateway.daily.reference.import.line"].create(line_values)
        next_action = self.env["ir.actions.actions"]._for_xml_id(
            "pagoflex_wallet_gateway.action_pf_gateway_daily_reference_import"
        )
        next_action.update(
            {
                "res_id": import_record.id,
                "view_mode": "form",
                "views": [(False, "form")],
                "target": "current",
            }
        )

        return {
            "type": "ir.actions.client",
            "tag": "display_notification",
            "params": {
                "title": _("Listado de Referencia Diario importado"),
                "message": _("Se importaron %s filas.") % len(line_values),
                "type": "success",
                "sticky": False,
                "next": next_action,
            },
        }
